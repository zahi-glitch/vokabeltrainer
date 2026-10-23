#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Vokabeltrainer - Die neue Linie 1 B1  (Fenster-Version / GUI)

Liest die drei Arbeitsmappen mit je einem Blatt pro Kapitel:
    B1_Verben.xlsx / B1_Nomen.xlsx / B1_Adjektive.xlsx
Blattname wie "K02 Hier kaufe ich ein" -> Kapitel 2.

Start:
    python trainer_gui.py
    python trainer_gui.py B1_Verben.xlsx B1_Nomen.xlsx ...
"""

import os
import re
import sys
import random

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    from openpyxl import load_workbook
except ImportError:
    _r = tk.Tk(); _r.withdraw()
    messagebox.showerror("Modul fehlt",
        "Das Modul 'openpyxl' fehlt.\n\nBitte im Terminal installieren:\n    pip install openpyxl")
    sys.exit(1)

DEFAULT_FILES = ["B1_Verben.xlsx", "B1_Nomen.xlsx", "B1_Adjektive.xlsx"]
TYPE_LABEL = {"verben": "Verb", "nomen": "Nomen", "adjektive": "Adjektiv"}
MENU_LABEL = {"verben": "Verben", "nomen": "Nomen", "adjektive": "Adjektive"}
ALL_CATEGORIES = "Alle"
ALL_KAP = "Alle Kapitel"
ALL_THEMES = "Alle Themen"
DASH = "—"

# Kapitel-Farben (hell): bg, fg
KAP_COLORS = [
    ("#DCEAF7", "#123048"), ("#E1EFD9", "#20361c"), ("#FBEFC6", "#4a3d0e"),
    ("#FBE1D3", "#4a2a1a"), ("#E7E0F5", "#2f2450"), ("#D8F0EC", "#103a33"),
    ("#FBDCE7", "#4a1730"), ("#E3ECF0", "#1d3540"),
]
DEFAULT_BG = "#F2F2F2"


# ---------------------------------------------------------------- Excel lesen
def script_dir():
    return os.path.dirname(os.path.abspath(__file__))


def find_default(name):
    for base in (script_dir(), os.getcwd()):
        p = os.path.join(base, name)
        if os.path.exists(p):
            return p
    return None


def cell(r, i):
    if i is None or i >= len(r) or r[i] is None:
        return ""
    return str(r[i]).strip()


def make_finders(header):
    def col(*names):
        for name in names:
            for i, h in enumerate(header):
                if h == name.lower():
                    return i
        return None

    def col_sw(prefix):
        for i, h in enumerate(header):
            if h.startswith(prefix):
                return i
        return None
    return col, col_sw


def detect_type(header):
    hs = set(header)
    if "infinitiv" in hs or any(h.startswith("partizip") for h in header):
        return "verben"
    if "artikel" in hs and any(h.startswith("nomen") for h in header):
        return "nomen"
    if "adjektiv" in hs or "komparativ" in hs or "superlativ" in hs:
        return "adjektive"
    return None


def parse_kapitel(sheet_name, index):
    m = re.match(r"^K(\d+)\s*(.*)$", sheet_name.strip())
    if m:
        return int(m.group(1)), (m.group(2).strip() or f"Kapitel {int(m.group(1))}")
    return index, sheet_name.strip()


def cards_from_sheet(ws, typ, kap_nr, kap_titel):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    col, col_sw = make_finders(header)
    idx_thema = col("thema", "tema", "topic", "theme")
    idx_en = col("english", "englisch", "inglese")
    out = []
    if typ == "verben":
        i_inf = col("infinitiv"); i_pr = col("präteritum", "praeteritum")
        i_pa = col("partizip ii", "partizip 2", "partizip"); i_vw = col_sw("verwendung")
        for r in rows[1:]:
            if not cell(r, i_inf):
                continue
            forms = " – ".join([x for x in (cell(r, i_inf), cell(r, i_pr), cell(r, i_pa)) if x and x != DASH])
            sol = [forms] + ([cell(r, i_vw)] if cell(r, i_vw) else [])
            out.append(dict(type=typ, kap=kap_nr, kaptitel=kap_titel, thema=cell(r, idx_thema),
                            prompt=cell(r, idx_en) or cell(r, i_inf), solution=sol))
    elif typ == "nomen":
        i_ar = col("artikel"); i_no = col_sw("nomen") or col("substantiv")
        i_pl = col("plural"); i_bs = col_sw("beispiel"); i_hi = col_sw("hinweis")
        for r in rows[1:]:
            if not cell(r, i_no):
                continue
            art = cell(r, i_ar); head = (art + " " if art else "") + cell(r, i_no)
            if cell(r, i_pl) and cell(r, i_pl) != DASH:
                head += ", " + cell(r, i_pl)
            sol = [head] + ([cell(r, i_bs)] if cell(r, i_bs) else [])
            if cell(r, i_hi):
                sol.append("Hinweis: " + cell(r, i_hi))
            out.append(dict(type=typ, kap=kap_nr, kaptitel=kap_titel, thema=cell(r, idx_thema),
                            prompt=cell(r, idx_en) or cell(r, i_no), solution=sol))
    elif typ == "adjektive":
        i_aj = col("adjektiv", "adjective"); i_ko = col("komparativ"); i_su = col("superlativ")
        i_ge = col_sw("gegenteil"); i_bs = col_sw("beispiel"); i_hi = col_sw("hinweis")
        for r in rows[1:]:
            if not cell(r, i_aj):
                continue
            parts = [cell(r, i_aj)]
            for x in (cell(r, i_ko), cell(r, i_su)):
                if x and x != DASH:
                    parts.append(x)
            sol = [" – ".join(parts)]
            if cell(r, i_ge) and cell(r, i_ge) != DASH:
                sol.append("Gegenteil: " + cell(r, i_ge))
            if cell(r, i_bs):
                sol.append(cell(r, i_bs))
            if cell(r, i_hi):
                sol.append("Hinweis: " + cell(r, i_hi))
            out.append(dict(type=typ, kap=kap_nr, kaptitel=kap_titel, thema=cell(r, idx_thema),
                            prompt=cell(r, idx_en) or cell(r, i_aj), solution=sol))
    return out


def load_all(paths):
    cards, problems = [], []
    for p in paths:
        if not os.path.exists(p):
            problems.append(f"Nicht gefunden: {os.path.basename(p)}"); continue
        try:
            wb = load_workbook(p, data_only=True)
        except Exception as e:
            problems.append(f"Fehler bei {os.path.basename(p)}: {e}"); continue
        for idx, ws in enumerate(wb.worksheets, start=1):
            head = [str(h).strip().lower() if h is not None else "" for h in
                    (next(ws.iter_rows(values_only=True), []) or [])]
            typ = detect_type(head)
            if not typ:
                continue
            nr, titel = parse_kapitel(ws.title, idx)
            cards.extend(cards_from_sheet(ws, typ, nr, titel))
    return cards, problems


# ---------------------------------------------------------------- Deck
class Deck:
    def __init__(self, cards, no_repeat=True):
        self.cards = cards; self.no_repeat = no_repeat
        self.pool = []; self.round = 0; self.pos = 0

    def draw(self):
        if not self.cards:
            return None, None, None
        if not self.no_repeat:
            return random.choice(self.cards), None, None
        if not self.pool:
            self.pool = self.cards[:]; random.shuffle(self.pool); self.round += 1; self.pos = 0
        self.pos += 1
        return self.pool.pop(), self.pos, len(self.cards)


# ---------------------------------------------------------------- GUI
class App:
    def __init__(self, root, cards):
        self.root = root
        self.cards = cards
        self.deck = None
        self.state = "prompt"
        self.current = None

        self.kapitel = tk.StringVar(value=ALL_KAP)
        self.wortart = tk.StringVar(value=ALL_CATEGORIES)
        self.thema = tk.StringVar(value=ALL_THEMES)
        self.no_repeat = tk.BooleanVar(value=True)

        root.title("Vokabeltrainer – Die neue Linie 1 B1")
        root.geometry("640x560"); root.minsize(540, 480); root.configure(bg="#FFFFFF")

        self._build()
        root.bind("<space>", lambda e: self.on_primary())
        root.bind("<Return>", lambda e: self.on_primary())
        root.bind("<Right>", lambda e: self.next_card())
        root.bind("q", lambda e: root.destroy())
        root.bind("<Escape>", lambda e: root.destroy())

        self.rebuild_deck()

    def _build(self):
        bar = ttk.Frame(self.root, padding=(12, 10)); bar.pack(fill="x")

        ttk.Label(bar, text="Kapitel:").grid(row=0, column=0, sticky="w")
        self.cb_kap = ttk.Combobox(bar, textvariable=self.kapitel, state="readonly", width=28)
        self.cb_kap.grid(row=0, column=1, columnspan=3, sticky="we", padx=(6, 0), pady=(0, 6))
        self.cb_kap.bind("<<ComboboxSelected>>", self.on_kapitel_change)

        ttk.Label(bar, text="Wortart:").grid(row=1, column=0, sticky="w")
        self.cb_wort = ttk.Combobox(bar, textvariable=self.wortart, state="readonly", width=14)
        self.cb_wort.grid(row=1, column=1, sticky="w", padx=(6, 12), pady=(0, 6))
        self.cb_wort.bind("<<ComboboxSelected>>", self.on_wortart_change)

        ttk.Label(bar, text="Thema:").grid(row=1, column=2, sticky="e")
        self.cb_thema = ttk.Combobox(bar, textvariable=self.thema, state="readonly", width=22)
        self.cb_thema.grid(row=1, column=3, sticky="we", padx=(6, 0), pady=(0, 6))
        self.cb_thema.bind("<<ComboboxSelected>>", lambda e: self.rebuild_deck())

        ttk.Checkbutton(bar, text="ohne Wiederholung", variable=self.no_repeat,
                        command=self.rebuild_deck).grid(row=2, column=0, columnspan=2, sticky="w")
        ttk.Button(bar, text="Dateien…", command=self.pick_files).grid(row=2, column=3, sticky="e")
        bar.columnconfigure(3, weight=1)

        outer = tk.Frame(self.root, bg="#FFFFFF"); outer.pack(fill="both", expand=True, padx=16, pady=6)
        self.card = tk.Frame(outer, bg=DEFAULT_BG, highlightthickness=1, highlightbackground="#D0D0D0")
        self.card.pack(fill="both", expand=True)

        self.lbl_eye = tk.Label(self.card, bg=DEFAULT_BG, fg="#555", font=("Helvetica", 11, "bold"))
        self.lbl_eye.pack(anchor="w", padx=18, pady=(14, 0))
        self.lbl_count = tk.Label(self.card, bg=DEFAULT_BG, fg="#888", font=("Helvetica", 10))
        self.lbl_count.place(relx=1.0, x=-16, y=14, anchor="ne")
        self.lbl_tag = tk.Label(self.card, bg=DEFAULT_BG, fg="#666", font=("Helvetica", 10, "bold"))
        self.lbl_tag.pack(anchor="w", padx=18, pady=(2, 0))
        self.lbl_prompt = tk.Label(self.card, bg=DEFAULT_BG, fg="#111",
                                   font=("Helvetica", 30, "bold"), wraplength=520, justify="center")
        self.lbl_prompt.pack(expand=True, pady=(4, 4))
        self.lbl_sol = tk.Label(self.card, bg=DEFAULT_BG, fg="#1F3D5A",
                                font=("Helvetica", 15), wraplength=520, justify="center")
        self.lbl_sol.pack(expand=True, pady=(0, 18), padx=18)
        self.card.bind("<Configure>", self._resize)

        foot = ttk.Frame(self.root, padding=(12, 10)); foot.pack(fill="x")
        self.btn = ttk.Button(foot, text="Lösung zeigen", command=self.on_primary)
        self.btn.pack(side="left", ipadx=8, ipady=4)
        ttk.Button(foot, text="Nächste →", command=self.next_card).pack(side="left", padx=8, ipady=4)
        ttk.Label(foot, text="Leertaste = weiter · q = beenden",
                  foreground="#999").pack(side="right")

        self.refresh_kapitel_options()
        self.refresh_wortart_options()

    # -------- options
    def _resize(self, e):
        w = max(300, e.width - 60)
        self.lbl_prompt.config(wraplength=w); self.lbl_sol.config(wraplength=w)

    def kapitel_list(self):
        return sorted({(c["kap"], c["kaptitel"]) for c in self.cards})

    def refresh_kapitel_options(self):
        vals = [ALL_KAP] + [f"K{nr} · {titel}" for nr, titel in self.kapitel_list()]
        self.cb_kap["values"] = vals
        if self.kapitel.get() not in vals:
            self.kapitel.set(ALL_KAP)

    def refresh_wortart_options(self):
        present = [t for t in ("verben", "nomen", "adjektive") if any(c["type"] == t for c in self.cards)]
        vals = [MENU_LABEL[t] for t in present]
        if len(present) > 1:
            vals.append(ALL_CATEGORIES)
        self.cb_wort["values"] = vals
        if self.wortart.get() not in vals:
            self.wortart.set(vals[-1] if vals else "")
        self.refresh_thema_options()

    def selected_kap_nr(self):
        v = self.kapitel.get()
        if v == ALL_KAP:
            return None
        m = re.match(r"^K(\d+)", v)
        return int(m.group(1)) if m else None

    def selected_types(self):
        w = self.wortart.get()
        if w == ALL_CATEGORIES:
            return ["verben", "nomen", "adjektive"]
        for t, lbl in MENU_LABEL.items():
            if lbl == w:
                return [t]
        return ["verben", "nomen", "adjektive"]

    def base_cards(self):
        knr = self.selected_kap_nr()
        types = self.selected_types()
        return [c for c in self.cards if (knr is None or c["kap"] == knr) and c["type"] in types]

    def refresh_thema_options(self):
        if self.selected_kap_nr() is None:
            self.cb_thema["values"] = [ALL_THEMES]
            self.thema.set(ALL_THEMES)
            self.cb_thema.configure(state="disabled")
            return
        self.cb_thema.configure(state="readonly")
        order = []
        for c in self.base_cards():
            if c["thema"] and c["thema"] not in order:
                order.append(c["thema"])
        vals = [ALL_THEMES] + order
        self.cb_thema["values"] = vals
        if self.thema.get() not in vals:
            self.thema.set(ALL_THEMES)

    def on_kapitel_change(self, e=None):
        self.thema.set(ALL_THEMES); self.refresh_thema_options(); self.rebuild_deck()

    def on_wortart_change(self, e=None):
        self.refresh_thema_options(); self.rebuild_deck()

    def selected_cards(self):
        c = self.base_cards()
        if self.selected_kap_nr() is not None and self.thema.get() != ALL_THEMES:
            c = [x for x in c if x["thema"] == self.thema.get()]
        return c

    def rebuild_deck(self):
        self.deck = Deck(self.selected_cards(), no_repeat=self.no_repeat.get())
        self.next_card()

    def _set_bg(self, bg, fg):
        for w in (self.card, self.lbl_eye, self.lbl_count, self.lbl_tag, self.lbl_prompt, self.lbl_sol):
            w.config(bg=bg)
        self.lbl_prompt.config(fg=fg)
        self.lbl_eye.config(fg=fg); self.lbl_tag.config(fg=fg)
        self.lbl_count.config(fg=fg); self.lbl_sol.config(fg=fg)

    def next_card(self):
        card, pos, total = self.deck.draw()
        self.current = card; self.state = "prompt"
        if card is None:
            self._set_bg(DEFAULT_BG, "#111")
            self.lbl_eye.config(text=""); self.lbl_tag.config(text=""); self.lbl_count.config(text="")
            self.lbl_prompt.config(text="Keine Vokabeln für diese Auswahl.")
            self.lbl_sol.config(text=""); self.btn.config(text="Lösung zeigen", state="disabled")
            return
        self.btn.config(state="normal", text="Lösung zeigen")
        bg, fg = KAP_COLORS[(card["kap"] - 1) % len(KAP_COLORS)]
        self._set_bg(bg, fg)
        self.lbl_eye.config(text=f"K{card['kap']} · {card['kaptitel']}")
        mixed = len(self.selected_types()) > 1
        self.lbl_tag.config(text=(f"[{TYPE_LABEL[card['type']]}]" if mixed else ""))
        if self.no_repeat.get() and pos is not None:
            extra = f"  ·  {self.deck.round}. Runde" if self.deck.round > 1 else ""
            self.lbl_count.config(text=f"{pos}/{total}{extra}")
        else:
            self.lbl_count.config(text="")
        self.lbl_prompt.config(text=card["prompt"])
        self.lbl_sol.config(text="")

    def reveal(self):
        if not self.current:
            return
        self.lbl_sol.config(text="\n".join(self.current["solution"]))
        self.state = "solution"; self.btn.config(text="Nächste Vokabel →")

    def on_primary(self):
        if self.state == "prompt":
            self.reveal()
        else:
            self.next_card()

    def pick_files(self):
        paths = filedialog.askopenfilenames(title="Vokabel-Dateien wählen",
                    filetypes=[("Excel", "*.xlsx"), ("Alle Dateien", "*.*")])
        if not paths:
            return
        cards, problems = load_all(list(paths))
        if not cards:
            messagebox.showerror("Keine Daten", "Keine gültigen Vokabeldateien.\n\n" + "\n".join(problems))
            return
        self.cards = cards
        if problems:
            messagebox.showwarning("Hinweis", "\n".join(problems))
        self.kapitel.set(ALL_KAP); self.wortart.set(ALL_CATEGORIES); self.thema.set(ALL_THEMES)
        self.refresh_kapitel_options(); self.refresh_wortart_options(); self.rebuild_deck()


def main():
    paths = sys.argv[1:] if len(sys.argv) > 1 else [p for p in (find_default(n) for n in DEFAULT_FILES) if p]
    root = tk.Tk()
    cards, problems = load_all(paths) if paths else ([], [])
    if not cards:
        root.withdraw()
        messagebox.showinfo("Dateien wählen",
            "Keine Vokabeldateien gefunden.\nBitte die drei Excel-Dateien auswählen.")
        chosen = filedialog.askopenfilenames(title="Vokabel-Dateien wählen",
                    filetypes=[("Excel", "*.xlsx"), ("Alle Dateien", "*.*")])
        cards, problems = load_all(list(chosen)) if chosen else ([], [])
        if not cards:
            messagebox.showerror("Keine Daten", "Keine gültigen Vokabeldateien gefunden.")
            root.destroy(); return
        root.deiconify()
    App(root, cards)
    root.mainloop()


if __name__ == "__main__":
    main()
